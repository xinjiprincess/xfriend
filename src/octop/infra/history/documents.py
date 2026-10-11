"""Select and render conversation sections without mutating conversation state."""

from __future__ import annotations

import asyncio
import hashlib
import html
import io
import json
import os
import sys
import zipfile
from pathlib import Path
from typing import Any

from octop.i18n import tr


def normalize_messages(messages: list[Any]) -> list[dict[str, Any]]:
    """Accept archive LangChain wire messages and already normalized engine messages."""
    result = []
    for message in messages:
        data = message.get("data") if isinstance(message.get("data"), dict) else message
        role = str(message.get("type") or data.get("role") or "assistant")
        role = {"human": "user", "ai": "assistant"}.get(role, role)
        normalized = {"role": role, "content": data.get("content", data.get("text", ""))}
        if data.get("agent_id"):
            normalized["agent_id"] = data["agent_id"]
        result.append(normalized)
    return result


def role_label(message: dict[str, Any], locale: str = "zh") -> str:
    agent = message.get("agent_id")
    if agent:
        return str(agent)
    role = str(message.get("role") or "assistant")
    role = "user" if role == "human" else "assistant" if role == "ai" else role
    return (
        tr(f"documents.roles.{role}", locale)
        if role in {"user", "assistant", "tool", "system"}
        else role
    )


def history_digest(messages: list[Any]) -> str:
    return hashlib.sha256(
        json.dumps(messages, sort_keys=True, ensure_ascii=False, default=str).encode()
    ).hexdigest()


def select_messages(
    messages: list[Any], mode: str, indices: list[int], start: int, end: int
) -> list[Any]:
    if not messages:
        raise ValueError("Empty history")
    if mode == "all":
        return messages
    if mode == "selected":
        chosen = sorted(set(indices))
    elif mode == "range":
        if start > end or start < 0 or end >= len(messages):
            raise ValueError("Invalid range")
        chosen = list(range(start, end + 1))
    elif mode == "turn":
        if (
            start < 0
            or start >= len(messages)
            or messages[start].get("role") not in {"user", "human"}
        ):
            raise ValueError("A turn must start with a user message")
        stop = next(
            (
                i
                for i in range(start + 1, len(messages))
                if messages[i].get("role") in {"user", "human"}
            ),
            len(messages),
        )
        chosen = list(range(start, stop))
    else:
        raise ValueError("Invalid selection mode")
    if not chosen or min(chosen) < 0 or max(chosen) >= len(messages):
        raise ValueError("Invalid message selection")
    return [messages[i] for i in chosen]


def message_text(message: dict[str, Any]) -> str:
    content = message.get("content", message.get("text", ""))
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "\n".join(
            block if isinstance(block, str) else str(block.get("text", ""))
            for block in content
            if isinstance(block, str | dict)
        )
    return json.dumps(content, ensure_ascii=False, default=str)


def document_html(messages: list[Any], title: str, locale: str = "zh") -> str:
    from markdown_it import MarkdownIt

    markdown = MarkdownIt("commonmark", {"html": False}).enable("table")
    markdown.disable("image")
    sections = []
    for message in messages:
        role = html.escape(role_label(message, locale))
        sections.append(
            f"<section><h2>{role}</h2>{markdown.render(message_text(message))}</section>"
        )
    return (
        '<!doctype html><html lang="zh"><head><meta charset="utf-8">'
        "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; style-src 'unsafe-inline'\">"
        f"<title>{html.escape(title)}</title><style>"
        "body{font:16px/1.7 'Microsoft YaHei',sans-serif;max-width:900px;margin:40px auto;padding:24px;color:#172033}"
        "h1{font-size:26px}h2{font-size:16px;color:#3757bd}section{border-top:1px solid #ddd;padding:12px 0}"
        "pre{white-space:pre-wrap;background:#f3f5fa;padding:12px}p,li{overflow-wrap:anywhere}"
        "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:6px}"
        f"</style></head><body><h1>{html.escape(title)}</h1>{''.join(sections)}</body></html>"
    )


def document_docx(messages: list[Any], title: str, locale: str = "zh") -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from markdown_it import MarkdownIt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading(title, 0)
    for message in messages:
        document.add_heading(role_label(message, locale), 2)
        tokens = (
            MarkdownIt("commonmark", {"html": False}).enable("table").parse(message_text(message))
        )
        paragraph: Any = None
        list_style: list[str] = []
        table_rows: list[list[str]] | None = None
        for token in tokens:
            if token.type == "table_open":
                table_rows = []
            elif token.type == "table_close" and table_rows:
                table = document.add_table(rows=0, cols=max(map(len, table_rows)))
                table.style = "Table Grid"
                for values in table_rows:
                    cells = table.add_row().cells
                    for index, value in enumerate(values):
                        cells[index].text = value
                table_rows = None
            elif table_rows is not None:
                if token.type == "tr_open":
                    table_rows.append([])
                elif token.type == "inline":
                    table_rows[-1].append(token.content)
            elif token.type in {"bullet_list_open", "ordered_list_open"}:
                list_style.append(
                    "List Bullet" if token.type == "bullet_list_open" else "List Number"
                )
            elif token.type in {"bullet_list_close", "ordered_list_close"}:
                list_style.pop()
            elif token.type == "heading_open":
                paragraph = document.add_heading(level=min(6, int(token.tag[1:])))
            elif token.type == "paragraph_open":
                paragraph = document.add_paragraph(style=list_style[-1] if list_style else None)
            elif token.type in {"fence", "code_block"}:
                paragraph = document.add_paragraph()
                paragraph.add_run(token.content.rstrip()).font.name = "Consolas"
            elif token.type == "inline":
                if paragraph is None:
                    paragraph = document.add_paragraph()
                bold = italic = False
                for child in token.children or []:
                    if child.type == "strong_open":
                        bold = True
                    elif child.type == "strong_close":
                        bold = False
                    elif child.type == "em_open":
                        italic = True
                    elif child.type == "em_close":
                        italic = False
                    elif child.type in {"text", "code_inline", "softbreak", "hardbreak"}:
                        run = paragraph.add_run(
                            "\n" if child.type.endswith("break") else child.content
                        )
                        run.bold, run.italic = bold, italic
                        if child.type == "code_inline":
                            run.font.name = "Consolas"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def bundled_browser() -> str | None:
    root = Path(sys.executable).resolve().parent.parent / "browsers"
    if not root.is_dir():
        return None
    names = {
        "chrome-headless-shell.exe",
        "headless_shell.exe",
        "chrome-headless-shell",
        "headless_shell",
    }
    return next(
        (str(path) for path in root.rglob("*") if path.name in names and path.is_file()), None
    )


async def render_document(
    messages: list[Any], title: str, format: str, locale: str = "zh"
) -> tuple[bytes, str, str]:
    if format == "docx":
        data = await asyncio.get_running_loop().run_in_executor(
            None, document_docx, messages, title, locale
        )
        return (
            data,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    markup = await asyncio.get_running_loop().run_in_executor(
        None, document_html, messages, title, locale
    )
    if format == "html":
        return markup.encode("utf-8"), "text/html; charset=utf-8", "html"
    if format not in {"pdf", "png"}:
        raise ValueError("Invalid document format")
    from playwright.async_api import async_playwright

    executable = await asyncio.get_running_loop().run_in_executor(None, bundled_browser)
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(
            executable_path=executable,
            channel="msedge" if os.name == "nt" and executable is None else None,
        )
        try:
            page = await browser.new_page(
                viewport={"width": 1100, "height": 1600}, device_scale_factor=1
            )
            await page.route("**/*", lambda route: route.abort())
            await page.set_content(markup, wait_until="load")
            if format == "pdf":
                return await page.pdf(format="A4", print_background=True), "application/pdf", "pdf"
            height = int(
                await page.evaluate("Math.ceil(document.body.getBoundingClientRect().bottom + 40)")
            )
            archive = io.BytesIO()
            with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as output:
                for index, top in enumerate(range(0, height, 1600)):
                    picture = await page.screenshot(
                        clip={"x": 0, "y": top, "width": 1100, "height": min(1600, height - top)}
                    )
                    output.writestr(f"conversation-{index + 1:03}.png", picture)
            return archive.getvalue(), "application/zip", "zip"
        finally:
            await browser.close()
