"""Section boundaries, stale snapshots, escaping, and editable DOCX structure."""

import io
from typing import Any

import pytest
from docx import Document

from octop.infra.history.documents import (
    bundled_browser,
    document_docx,
    document_html,
    history_digest,
    normalize_messages,
    select_messages,
)

MESSAGES: list[Any] = [
    {"role": "user", "content": "First"},
    {"role": "assistant", "content": "Answer"},
    {"role": "tool", "content": "Result"},
    {"role": "user", "content": "Second"},
    {"role": "assistant", "content": "Another"},
]


def test_archive_wire_messages_are_normalized_before_selection() -> None:
    messages = normalize_messages(
        [
            {"type": "human", "data": {"content": "中文问题", "id": "u"}},
            {"type": "ai", "data": {"content": [{"type": "text", "text": "中文回答"}]}},
        ]
    )
    assert messages[0]["role"] == "user"
    selected = select_messages(messages, "turn", [], 0, 0)
    html = document_html(selected, "对话")
    assert "中文问题" in html and "中文回答" in html


def test_portable_export_uses_its_bundled_browser(tmp_path, monkeypatch) -> None:
    import sys

    root = tmp_path / "portable"
    executable = (
        root
        / "browsers"
        / "chromium_headless_shell-1228"
        / "chrome-headless-shell-win64"
        / "chrome-headless-shell.exe"
    )
    executable.parent.mkdir(parents=True)
    executable.touch()
    monkeypatch.setattr(sys, "executable", str(root / "runtime" / "python.exe"))
    assert bundled_browser() == str(executable)


def test_selection_boundaries_and_order() -> None:
    assert select_messages(MESSAGES, "turn", [], 0, 0) == MESSAGES[:3]
    assert select_messages(MESSAGES, "turn", [], 3, 0) == MESSAGES[3:]
    assert select_messages(MESSAGES, "selected", [4, 0, 4], 0, 0) == [MESSAGES[0], MESSAGES[4]]
    assert select_messages(MESSAGES, "range", [], 1, 3) == MESSAGES[1:4]
    assert select_messages(MESSAGES, "all", [], 0, 0) == MESSAGES
    with pytest.raises(ValueError):
        select_messages(MESSAGES, "turn", [], 1, 0)
    with pytest.raises(ValueError):
        select_messages(MESSAGES, "range", [], 1, 10**10)
    with pytest.raises(ValueError):
        select_messages(MESSAGES, "selected", [-1], 0, 0)


def test_digest_changes_and_html_is_inert() -> None:
    assert history_digest(MESSAGES) != history_digest(
        MESSAGES + [{"role": "user", "content": "New"}]
    )
    markup = document_html(
        [
            {
                "role": "user",
                "content": "<script>alert(1)</script>\n![remote](https://example.com/pixel)",
            }
        ],
        "<b>title</b>",
    )
    assert "<script>" not in markup
    assert "<img" not in markup
    assert "&lt;b&gt;title&lt;/b&gt;" in markup
    assert "Content-Security-Policy" in markup
    assert "你" in markup


def test_docx_preserves_chinese_bold_and_table() -> None:
    data = document_docx(
        [
            {
                "role": "assistant",
                "content": "# 标题\n中文 **加粗**\n\n| A | B |\n| --- | --- |\n| 甲 | 乙 |",
            }
        ],
        "对话",
    )
    document = Document(io.BytesIO(data))
    assert any(paragraph.text == "标题" for paragraph in document.paragraphs)
    assert any(run.text == "加粗" and run.bold for p in document.paragraphs for run in p.runs)
    assert document.tables[0].cell(1, 0).text == "甲"
